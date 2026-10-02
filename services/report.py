"""Admin Excel report: one branded workbook with an Overview (KPI tiles, charts,
summary tables) plus a sheet per inbox, visitors and the Sellsy pipeline.

build_workbook() takes the same shaped records the dashboard renders, so it
never queries anything itself (routes/admin.py export_report does one round
trip). Storage plumbing (file URLs/paths) is left out, as in the e-mails.
"""

from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.chart import BarChart, DoughnutChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.formatting.rule import DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from content.admin_mock import brief_value
from services.mailer import _human_size
from services.store import AST

NAVY, INK, GOLD, GOLD_DEEP = "161F3B", "0E1528", "C9A15B", "B8822F"
IVORY, SAND, BAND, LINE = "FBF7F0", "F4EDE1", "F8F4EC", "E6DCCB"
TEAL, MUTED, WHITE = "1F8C87", "6E7682", "FFFFFF"
SERIES = ["1F8C87", "C9A15B", "4A78A8", "8E6FB0", "D97B5E", "6C8F4E"]

# (fill, text) per status, matching the admin's chip colours.
STATUS_COLORS = {
    "New": ("E1F2F0", "1F6F6B"), "New Lead": ("E1F2F0", "1F6F6B"), "Open": ("E1F2F0", "1F6F6B"),
    "Contacted": ("FBF0DA", "8A6420"), "Returning": ("FBF0DA", "8A6420"), "Late": ("FBF0DA", "8A6420"),
    "Qualified": ("E4EEF8", "2E5C8A"), "Closed": ("E4EEF8", "2E5C8A"),
    "Won": ("E2F3E5", "2E7D44"), "Converted": ("E2F3E5", "2E7D44"),
    "Lost": ("FBE5E0", "A2452F"),
}

SHORT_NAMES = {"British Virgin Islands": "BVI"}
FUNNEL_LABELS = {"Engaged with a module": "Engaged", "Used the Estimator": "Used estimator",
                 "Submitted a Brief": "Sent a brief"}

USD, EUR, INT = '"$"#,##0', '"€"#,##0', "#,##0"
SERIF = "Georgia"

thin = Side(style="thin", color=LINE)
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)


def _fill(color):
    return PatternFill("solid", start_color=color, end_color=color)


def _put(ws, row, col, value, **style):
    """Write a cell; text is always text (a form value starting with '=' must
    never become a live formula in the admin's Excel)."""
    cell = ws.cell(row=row, column=col, value=value)
    if isinstance(value, str):
        cell.data_type = "s"
    for key, val in style.items():
        setattr(cell, key, val)
    return cell


def _iso_date(text):
    try:
        return date.fromisoformat(str(text)[:10])
    except ValueError:
        return None


def _pct(part, whole):
    return part / whole if whole else 0


def _day(d):
    """"2 Oct" without strftime('%-d'), which Windows does not support."""
    return f"{d.day} {d:%b}"


def _stamp(dt):
    return f"{_day(dt)} {dt:%Y, %H:%M}"


# ---------------------------------------------------------------- overview
def _section(ws, row, title):
    _put(ws, row, 2, title.upper(), font=Font(name="Calibri", size=10, bold=True, color=GOLD_DEEP))
    for col in range(2, 14):
        ws.cell(row=row, column=col).border = Border(bottom=Side(style="thin", color=GOLD))
    ws.row_dimensions[row].height = 22


def _mini_table(ws, row, col, headers, rows, formats=None):
    """Small styled table on the Overview; returns (first_data_row, last_data_row)."""
    for i, head in enumerate(headers):
        _put(ws, row, col + i, head, font=Font(bold=True, size=9, color=WHITE), fill=_fill(NAVY),
             alignment=Alignment(horizontal="left" if i == 0 else "right", vertical="center"), border=BOX)
    ws.row_dimensions[row].height = 20
    for r, values in enumerate(rows, start=1):
        for i, value in enumerate(values):
            cell = _put(ws, row + r, col + i, value, font=Font(size=10, color=INK), border=BOX,
                        fill=_fill(BAND if r % 2 == 0 else WHITE),
                        alignment=Alignment(horizontal="left" if i == 0 else "right", vertical="center"))
            if formats and formats[i]:
                cell.number_format = formats[i]
    if not rows:
        _put(ws, row + 1, col, "No data yet", font=Font(italic=True, size=9, color=MUTED))
    return row + 1, row + max(len(rows), 1)


def _kpi(ws, col, label, value, fmt, note):
    """A two-column KPI tile at rows 6-9."""
    span = f"{get_column_letter(col)}{{r}}:{get_column_letter(col + 1)}{{r}}"
    for r in (6, 7, 8, 9):
        ws.merge_cells(span.format(r=r))
        for c in (col, col + 1):
            ws.cell(row=r, column=c).fill = _fill(IVORY)
            ws.cell(row=r, column=c).border = Border(
                left=thin if c == col else None, right=thin if c == col + 1 else None,
                top=Side(style="thick", color=GOLD) if r == 6 else None, bottom=thin if r == 9 else None)
    _put(ws, 6, col, label.upper(), font=Font(size=8.5, bold=True, color=GOLD_DEEP),
         alignment=Alignment(horizontal="left", vertical="bottom", indent=1))
    cell = _put(ws, 7, col, value, font=Font(name=SERIF, size=24, color=NAVY),
                alignment=Alignment(horizontal="left", vertical="center", indent=1))
    if fmt:
        cell.number_format = fmt
    _put(ws, 8, col, note, font=Font(size=9, color=MUTED),
         alignment=Alignment(horizontal="left", vertical="top", indent=1, wrap_text=True))


def _line_chart(ws, daily, anchor):
    chart = LineChart()
    chart.title = "Sessions and event briefs per day"
    chart.height, chart.width = 7.4, 15.2
    chart.y_axis.title = None
    chart.x_axis.delete = chart.y_axis.delete = False  # openpyxl 3.1 hides axes by default
    chart.y_axis.number_format = "0"
    chart.x_axis.number_format = "d mmm"
    chart.x_axis.tickLblSkip = 2
    chart.y_axis.scaling.min = 0
    chart.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill=LINE))
    chart.legend.position, chart.legend.overlay = "b", False
    n = daily.max_row
    chart.add_data(Reference(daily, min_col=2, max_col=3, min_row=1, max_row=n), titles_from_data=True)
    chart.set_categories(Reference(daily, min_col=1, min_row=2, max_row=n))
    for series, color, width in zip(chart.series, (TEAL, GOLD), (28575, 22225)):
        series.graphicalProperties.line.solidFill = color
        series.graphicalProperties.line.width = width
        series.smooth = False  # counts: a smoothed curve dips below zero
    ws.add_chart(chart, anchor)


def _doughnut(ws, first, last, col, anchor):
    chart = DoughnutChart(holeSize=58)
    chart.height, chart.width = 7.4, 7.6
    chart.add_data(Reference(ws, min_col=col + 1, min_row=first - 1, max_row=last), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=col, min_row=first, max_row=last))
    series = chart.series[0]
    for i in range(last - first + 1):
        point = DataPoint(idx=i)
        point.graphicalProperties.solidFill = SERIES[i % len(SERIES)]
        point.graphicalProperties.line.solidFill = WHITE
        series.dPt.append(point)
    chart.legend.position, chart.legend.overlay = "b", False
    ws.add_chart(chart, anchor)


def _bar_chart(ws, first, last, col, anchor, color, horizontal=True):
    chart = BarChart()
    chart.type = "bar" if horizontal else "col"
    chart.height, chart.width = 7.2, 11.5
    chart.legend = None
    chart.y_axis.majorGridlines = None
    chart.x_axis.delete, chart.y_axis.delete = False, True  # categories shown; values are on the bars
    chart.add_data(Reference(ws, min_col=col + 1, min_row=first - 1, max_row=last), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=col, min_row=first, max_row=last))
    if horizontal:
        chart.x_axis.scaling.orientation = "maxMin"  # largest on top, matching the table
    series = chart.series[0]
    series.graphicalProperties.solidFill = color
    series.graphicalProperties.line.solidFill = color
    chart.gapWidth = 60
    series.dLbls = DataLabelList()
    series.dLbls.showVal = True
    series.dLbls.showSerName = series.dLbls.showCatName = series.dLbls.showLegendKey = False
    ws.add_chart(chart, anchor)


def _overview(wb, ctx):
    ws = wb.active
    ws.title = "Overview"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = GOLD
    ws.column_dimensions["A"].width = 2.5
    for col in range(2, 14):
        ws.column_dimensions[get_column_letter(col)].width = 13.5
    ws.column_dimensions["N"].width = 2.5

    s, days = ctx["summary"], ctx["days"]
    period = ""
    if days:
        first, last = date.fromisoformat(days[0]), date.fromisoformat(days[-1])
        period = f"{_day(first)} to {_day(last)} {last:%Y}"

    # Title band
    for r, h in ((2, 40), (3, 22), (4, 4)):
        ws.row_dimensions[r].height = h
    for r in (2, 3):
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=13)
        for c in range(2, 14):
            ws.cell(row=r, column=c).fill = _fill(NAVY)
    for c in range(2, 14):
        ws.cell(row=4, column=c).fill = _fill(GOLD)
    _put(ws, 2, 2, "Caribbean Incentive  ·  Command Centre Report",
         font=Font(name=SERIF, size=20, italic=True, color=WHITE),
         alignment=Alignment(vertical="center", indent=1))
    _put(ws, 3, 2, f"Last 14 days: {period}     ·     Generated {_stamp(ctx['generated'])} AST",
         font=Font(size=10, color="D9C9A3"), alignment=Alignment(vertical="center", indent=1))

    # KPI tiles
    sellsy = ctx.get("sellsy") or {}
    delta = s.get("sessions_delta")
    tiles = [
        ("Pipeline value", s["pipeline_value"], USD,
         f"${s['won_value']:,.0f} won · open and won briefs"),
        ("Sessions", s["sessions_14d"], INT,
         f"{s['sessions_week']} this week" + (f" · {delta:+.1f}% vs last" if isinstance(delta, (int, float)) else "")),
        ("Event briefs", s["total_briefs"], INT, f"{s['awaiting_reply']} awaiting first reply"),
        ("Lead conversion", _pct(s["converted"], s["total_visitors"]), "0%",
         f"{s['converted']} of {s['total_visitors']} visitors converted"),
        ("Visitors tracked", s["total_visitors"], INT, f"{s['engagement_avg']}% engaged with the site"),
        ("Sellsy pipeline", sellsy.get("open_value", 0), EUR,
         f"{sellsy.get('open_count', 0)} open · {sellsy.get('won_count', 0)} won" if sellsy else "Sellsy not available"),
    ]
    for i, (label, value, fmt, note) in enumerate(tiles):
        _kpi(ws, 2 + i * 2, label, value, fmt, note)
    ws.row_dimensions[6].height, ws.row_dimensions[7].height = 20, 38
    ws.row_dimensions[8].height, ws.row_dimensions[9].height = 28, 6

    # Traffic & demand
    _section(ws, 11, "Traffic and demand · last 14 days")
    _line_chart(ws, ctx["daily_ws"], "B13")
    sources = ctx["traffic_sources"]
    total_src = sum(x["visitors"] for x in sources)
    f, l = _mini_table(ws, 13, 11, ["Source", "Visitors", "Share"],
                       [(x["source"], x["visitors"], _pct(x["visitors"], total_src)) for x in sources],
                       [None, INT, "0%"])
    if sources:
        _doughnut(ws, f, l, 11, "H13")

    # Pipeline & conversion
    row = 30
    _section(ws, row, "Pipeline and conversion")
    stages = s["stages"]
    f, l = _mini_table(ws, row + 2, 2, ["Stage", "Briefs", "Value (USD)"],
                       [(x["stage"], x["count"], x["value"]) for x in stages], [None, INT, USD])
    ws.conditional_formatting.add(f"D{f}:D{l}", DataBarRule(start_type="num", start_value=0, end_type="max",
                                                             color=GOLD, showValue=True))
    funnel = ctx["conversion_funnel"]
    top = funnel[0]["count"] if funnel else 0
    f, l = _mini_table(ws, row + 2, 6, ["Conversion funnel", "Visitors", "% of visitors"],
                       [(FUNNEL_LABELS.get(x["stage"], x["stage"]), x["count"], _pct(x["count"], top)) for x in funnel],
                       [None, INT, "0%"])
    ws.conditional_formatting.add(f"H{f}:H{l}", DataBarRule(start_type="num", start_value=0, end_type="num",
                                                             end_value=1, color=TEAL, showValue=True))
    q = s["quick_stats"]
    ds = q["device_split"]
    _mini_table(ws, row + 2, 10, ["At a glance", "", "Value"], [
        ("Avg. session", "", q["avg_session_duration"]),
        ("Bounce rate", "", q["bounce_rate"] / 100),
        ("Top referrer", "", q["top_referrer"]),
        ("Peak hour (AST)", "", q["peak_traffic_hour"]),
        ("Desktop / mobile", "", f"{ds.get('Desktop', 0)}% / {ds.get('Mobile', 0)}%"),
        ("Site conversion", "", s["site_conversion"] / 100),
    ], [None, None, None])
    for r, fmt in ((row + 4, "0%"), (row + 8, "0.0%")):
        ws.cell(row=r, column=12).number_format = fmt
    for r in range(row + 2, row + 9):
        ws.merge_cells(start_row=r, start_column=10, end_row=r, end_column=11)

    # Interest
    row = 42
    _section(ws, row, "Where interest goes")
    islands = ctx["island_views"]
    f, l = _mini_table(ws, row + 2, 2, ["Island", "Page views"], [(SHORT_NAMES.get(x["island"], x["island"]), x["views"]) for x in islands],
                       [None, INT])
    if islands:
        _bar_chart(ws, f, l, 2, f"E{row + 2}", TEAL)
    modules = ctx["module_stats"]
    f, l = _mini_table(ws, row + 2, 10, ["Site module", "Views", "Engaged", "Avg. time"],
                       [(m["module"], m["views"], m["engagement_rate"] / 100, f"{m['avg_time_sec']}s")
                        for m in modules], [None, INT, "0%", None])
    ws.conditional_formatting.add(f"L{f}:L{l}", DataBarRule(start_type="num", start_value=0, end_type="num",
                                                             end_value=1, color=GOLD, showValue=True))

    _put(ws, row + 18, 2, "Figures come live from the website database and Sellsy at the moment of export. "
                          "Pipeline values use the midpoint of each brief's budget band.",
         font=Font(italic=True, size=9, color=MUTED))

    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.freeze_panes = None


def _daily_sheet(wb, ctx):
    ws = wb.create_sheet("Daily")
    ws.sheet_properties.tabColor = TEAL
    s = ctx["summary"]
    for i, head in enumerate(["Day", "Sessions", "Event briefs", "Pipeline value (USD)"], start=1):
        _put(ws, 1, i, head, font=Font(bold=True, color=WHITE, size=10), fill=_fill(NAVY),
             alignment=Alignment(horizontal="center", vertical="center"))
    for r, day in enumerate(ctx["days"], start=2):
        ws.cell(row=r, column=1, value=_iso_date(day)).number_format = "ddd d mmm yyyy"
        ws.cell(row=r, column=2, value=s["sessions_daily"][r - 2]).number_format = INT
        ws.cell(row=r, column=3, value=s["briefs_daily"][r - 2]).number_format = INT
        ws.cell(row=r, column=4, value=s["pipeline_daily"][r - 2]).number_format = USD
    for col, width in zip("ABCD", (20, 12, 14, 22)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    return ws


# ------------------------------------------------------------ data sheets
def _data_sheet(wb, title, subtitle, columns, records, tab, status_key=None):
    """columns: (header, getter, width, number_format, wrap)."""
    ws = wb.create_sheet(title)
    ws.sheet_properties.tabColor = tab
    ws.sheet_view.showGridLines = False
    last_col = get_column_letter(len(columns))

    ws.merge_cells(f"A1:{last_col}1")
    _put(ws, 1, 1, title, font=Font(name=SERIF, size=16, italic=True, color=NAVY),
         alignment=Alignment(vertical="center"))
    ws.merge_cells(f"A2:{last_col}2")
    _put(ws, 2, 1, subtitle, font=Font(size=9.5, color=MUTED), alignment=Alignment(vertical="top"))
    ws.row_dimensions[1].height, ws.row_dimensions[2].height = 28, 18

    for i, (head, _get, width, _fmt, _wrap) in enumerate(columns, start=1):
        _put(ws, 3, i, head, font=Font(bold=True, size=10, color=WHITE), fill=_fill(NAVY),
             alignment=Alignment(vertical="center", wrap_text=True), border=BOX)
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[3].height = 24

    for r, rec in enumerate(records, start=4):
        for i, (_head, get, _width, fmt, wrap) in enumerate(columns, start=1):
            value = get(rec)
            cell = _put(ws, r, i, value if value not in (None, "—") else "",
                        font=Font(size=10, color=INK), border=Border(bottom=thin),
                        alignment=Alignment(vertical="top", wrap_text=wrap))
            if fmt:
                cell.number_format = fmt

    end = 3 + max(len(records), 1)
    if not records:
        _put(ws, 4, 1, "Nothing recorded yet.", font=Font(italic=True, size=10, color=MUTED))
    else:
        rng = f"A4:{last_col}{end}"
        ws.conditional_formatting.add(rng, FormulaRule(formula=["MOD(ROW(),2)=1"], fill=_fill(BAND)))
        if status_key is not None:
            col = get_column_letter(status_key)
            for status, (bg, fg) in STATUS_COLORS.items():
                ws.conditional_formatting.add(
                    f"{col}4:{col}{end}",
                    FormulaRule(formula=[f'${col}4="{status}"'], fill=_fill(bg),
                                font=Font(bold=True, color=fg), stopIfTrue=True))
        ws.auto_filter.ref = f"A3:{last_col}{end}"
    ws.freeze_panes = "B4"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "3:3"
    return ws


def _subtitle(n, noun, generated):
    return f"{n} {noun}{'' if n == 1 else 's'}  ·  exported {_stamp(generated)} AST  ·  newest first"


def _yes(flag):
    return "Yes" if flag else "No"


def build_workbook(data):
    """ctx: summary, days, traffic_sources, conversion_funnel, island_views,
    module_stats (analytics.build), briefs (shape_briefs), rfps/callbacks
    (shape_inquiries), visitors (shape_visitors), sellsy (snapshot or None)."""
    ctx: dict = {**data, "generated": datetime.now(AST)}
    wb = Workbook()
    wb.properties.title = "Caribbean Incentive · Command Centre Report"
    wb.properties.creator = "Caribbean Incentive"

    daily = _daily_sheet(wb, ctx)
    ctx["daily_ws"] = daily
    _overview(wb, ctx)
    g = ctx["generated"]

    _data_sheet(wb, "Event Briefs", _subtitle(len(ctx["briefs"]), "event brief", g), [
        ("Reference", lambda b: b["reference"], 18, None, False),
        ("Submitted (AST)", lambda b: b["submitted_at"].replace(" AST", ""), 17, None, False),
        ("Status", lambda b: b["status"], 12, None, False),
        ("Company", lambda b: b["company"], 24, None, True),
        ("Contact", lambda b: b["contact_name"], 20, None, True),
        ("E-mail", lambda b: b["email"], 30, None, False),
        ("Phone", lambda b: b["phone"], 17, None, False),
        ("Country", lambda b: b["country"], 14, None, False),
        ("Request type", lambda b: b["request_type"], 20, None, True),
        ("Group size", lambda b: b["group_size"] or None, 11, INT, False),
        ("Preferred dates", lambda b: b["preferred_dates"], 18, None, True),
        ("Budget band", lambda b: b["budget"], 15, None, False),
        ("Pipeline value", lambda b: brief_value(b), 15, USD, False),
        ("About the event", lambda b: b["message"], 55, None, True),
    ], ctx["briefs"], GOLD, status_key=3)

    _data_sheet(wb, "RFP Uploads", _subtitle(len(ctx["rfps"]), "RFP", g), [
        ("Reference", lambda r: r.get("reference"), 18, None, False),
        ("Submitted (AST)", lambda r: (r.get("submitted_at") or "").replace(" AST", ""), 17, None, False),
        ("Status", lambda r: r.get("status") or "New", 12, None, False),
        ("Name", lambda r: r.get("name"), 22, None, True),
        ("Company", lambda r: r.get("company"), 24, None, True),
        ("E-mail", lambda r: r.get("email"), 30, None, False),
        ("Document", lambda r: r.get("file_name"), 36, None, True),
        ("Size", lambda r: _human_size(r.get("file_size")) if r.get("file_size") else "", 10, None, False),
    ], ctx["rfps"], "4A78A8", status_key=3)

    _data_sheet(wb, "Callbacks", _subtitle(len(ctx["callbacks"]), "callback request", g), [
        ("Reference", lambda r: r.get("reference"), 18, None, False),
        ("Submitted (AST)", lambda r: (r.get("submitted_at") or "").replace(" AST", ""), 17, None, False),
        ("Status", lambda r: r.get("status") or "New", 12, None, False),
        ("Name", lambda r: r.get("name"), 22, None, True),
        ("Company", lambda r: r.get("company"), 24, None, True),
        ("E-mail", lambda r: r.get("email"), 30, None, False),
        ("Phone", lambda r: r.get("phone"), 18, None, False),
    ], ctx["callbacks"], "8E6FB0", status_key=3)

    _data_sheet(wb, "Visitors", _subtitle(len(ctx["visitors"]), "visitor", g), [
        ("Visitor", lambda v: v["label"], 14, None, False),
        ("Status", lambda v: v["status"], 12, None, False),
        ("Company", lambda v: v.get("company"), 22, None, True),
        ("E-mail", lambda v: v.get("email"), 28, None, False),
        ("City", lambda v: "" if v["city"] == "Locating…" else v["city"], 20, None, True),
        ("Country", lambda v: v["country"], 14, None, False),
        ("Source", lambda v: v["source"], 14, None, False),
        ("Device", lambda v: v["device"], 10, None, False),
        ("Sessions", lambda v: v.get("sessions") or 0, 10, INT, False),
        ("Pages", lambda v: v.get("pages_viewed") or 0, 9, INT, False),
        ("Most viewed module", lambda v: v["top_module"], 20, None, False),
        ("First seen", lambda v: v["first_seen"], 18, None, False),
        ("Last seen", lambda v: v["last_seen"], 18, None, False),
        ("Converted via", lambda v: v.get("converted_reference") or "", 18, None, False),
    ], ctx["visitors"], TEAL, status_key=2)

    sellsy = ctx.get("sellsy")
    if sellsy:
        labels = {"open": "Open", "late": "Late", "won": "Won", "lost": "Lost", "closed": "Closed"}
        title = "Sellsy Pipeline"
        _data_sheet(wb, title, (f"{sellsy.get('pipeline') or 'Sellsy'}  ·  " +
                                _subtitle(len(sellsy["opportunities"]), "opportunity", g).replace("opportunitys", "opportunities")), [
            ("Number", lambda o: o["number"], 12, None, False),
            ("Opportunity", lambda o: o["name"], 34, None, True),
            ("Status", lambda o: labels.get(o["status"], o["status"].title()), 11, None, False),
            ("Stage", lambda o: o["step"], 24, None, True),
            ("Amount", lambda o: o["amount"], 14, EUR, False),
            ("Company", lambda o: o["company"], 24, None, True),
            ("Contact", lambda o: o["contact"], 20, None, True),
            ("E-mail", lambda o: o["email"], 28, None, False),
            ("Source", lambda o: o["source"], 24, None, True),
            ("From website", lambda o: _yes(o["is_website"]), 12, None, False),
            ("Created", lambda o: _iso_date(o["created"]), 13, "d mmm yyyy", False),
        ], sellsy["opportunities"], "D97B5E", status_key=3)

    wb.active = 0
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()

